from pathlib import Path
from openpyxl import load_workbook
from PIL import Image
import hashlib

BASE = Path(__file__).resolve().parent
XLSX = BASE / 'SISTEMA_PEDIDOS_SERGIO_V1.xlsx'
IMG_DIR = BASE / 'catalogo_imagenes'


def image_hash(path):
    with Image.open(path) as im:
        im = im.convert('RGB')
        im.thumbnail((32, 32))
        return hashlib.sha256(im.tobytes()).hexdigest()[:16]


def main():
    if not XLSX.exists():
        raise FileNotFoundError(f'No encuentro: {XLSX}')
    IMG_DIR.mkdir(exist_ok=True)
    wb = load_workbook(XLSX)
    if 'CATALOGO' not in wb.sheetnames:
        raise ValueError('No existe la hoja CATALOGO.')
    ws = wb['CATALOGO']
    headers = {c.value: c.column for c in ws[1] if c.value}
    for h in ('ID_PRODUCTO', 'IMAGEN', 'HASH_IMAGEN'):
        if h not in headers:
            raise ValueError(f'Falta la columna {h} en CATALOGO.')
    updated = missing = 0
    for row in range(2, ws.max_row + 1):
        product_id = ws.cell(row, headers['ID_PRODUCTO']).value
        if not product_id:
            continue
        image = next((IMG_DIR / f'{product_id}{ext}' for ext in ('.jpg','.jpeg','.png','.webp') if (IMG_DIR / f'{product_id}{ext}').exists()), None)
        if image is None:
            missing += 1
            continue
        ws.cell(row, headers['IMAGEN']).value = str(image.relative_to(BASE))
        ws.cell(row, headers['HASH_IMAGEN']).value = image_hash(image)
        updated += 1
    wb.save(XLSX)
    print(f'Catálogo actualizado. Con imagen: {updated}. Sin imagen: {missing}.')

if __name__ == '__main__':
    main()
